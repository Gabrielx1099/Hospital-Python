from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from flask_mail import Mail
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime, date, timedelta
from sqlalchemy import func, desc
from dotenv import load_dotenv
import os
import io
import joblib

load_dotenv()

from correo import enviar_correo_confirmacion
app = Flask(__name__)

# Carga tolerante a fallos del modelo de síntomas
modelo_sintomas = vectorizer_sintomas = None
try:
    if os.path.exists('models/vectorizer_sintomas.pkl') and os.path.exists('models/modelo_sintomas.pkl'):
        vectorizer_sintomas = joblib.load('models/vectorizer_sintomas.pkl')
        modelo_sintomas = joblib.load('models/modelo_sintomas.pkl')
        app.logger.info("✅ Modelo de clasificación de síntomas cargado correctamente")
    else:
        app.logger.warning("⚠️ Archivos del modelo no encontrados en 'models/'")
except Exception as e:
    app.logger.warning('Modelo de sintomas no disponible: %s', e)

def clasificar_sintomas(texto):
    if not modelo_sintomas or not vectorizer_sintomas:
        # Fallback a Medicina General
        med_gen = Especialidad.query.filter_by(nombre='Medicina General').first()
        return {
            'especialidad': 'Medicina General',
            'especialidad_id': med_gen.id if med_gen else None,
            'confianza': 0.0,
            'mensaje': 'Clasificador de síntomas fuera de línea temporalmente.'
        }
    
    # Vectorizar
    vec = vectorizer_sintomas.transform([texto])
    
    # Obtener probabilidades
    probs = modelo_sintomas.predict_proba(vec)[0]
    classes = modelo_sintomas.classes_
    
    # Encontrar clase con mayor probabilidad
    max_idx = probs.argmax()
    label = classes[max_idx]
    confianza = float(probs[max_idx])
    
    # Buscar especialidad en DB
    esp = Especialidad.query.filter_by(nombre=label).first()
    esp_id = esp.id if esp else None
    
    # Si la confianza es menor a 40%, devolver Medicina General como fallback
    if confianza < 0.40:
        med_gen = Especialidad.query.filter_by(nombre='Medicina General').first()
        return {
            'especialidad': 'Medicina General',
            'especialidad_id': med_gen.id if med_gen else None,
            'confianza': confianza,
            'mensaje': f'Confianza baja ({int(confianza*100)}%). Se recomienda Medicina General por seguridad.'
        }
        
    return {
        'especialidad': label,
        'especialidad_id': esp_id,
        'confianza': confianza
    }

app.config['SECRET_KEY'] = 'hospital_lanfranco_secret_2024'

app.config['SQLALCHEMY_DATABASE_URI'] = 'mysql+pymysql://root@localhost:3306/hospital_lanfranco'

app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

app.config['MAIL_SERVER'] = os.environ.get('MAIL_SERVER', 'smtp.gmail.com')
app.config['MAIL_PORT'] = int(os.environ.get('MAIL_PORT', 587))
app.config['MAIL_USE_TLS'] = os.environ.get('MAIL_USE_TLS', 'true').lower() == 'true'
app.config['MAIL_USERNAME'] = os.environ.get('MAIL_USERNAME')
app.config['MAIL_PASSWORD'] = os.environ.get('MAIL_PASSWORD')
app.config['MAIL_DEFAULT_SENDER'] = os.environ.get('MAIL_DEFAULT_SENDER', app.config['MAIL_USERNAME'])

db = SQLAlchemy(app)
mail = Mail(app)
try:
    with app.app_context():
        db.engine.connect()
        print("[OK] Conectado a MySQL correctamente")
except Exception as e:
    print("[ERROR] Error de conexion a MySQL:")
    print(e)

login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Por favor inicia sesion para acceder.'

class Usuario(UserMixin, db.Model):
    __tablename__ = 'usuarios'
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(100), nullable=False)
    apellido = db.Column(db.String(100), nullable=False)
    dni = db.Column(db.String(8), unique=True, nullable=False)
    email = db.Column(db.String(150), unique=True, nullable=False)
    telefono = db.Column(db.String(15))
    password = db.Column(db.String(200), nullable=False)
    rol = db.Column(db.String(20), default='paciente')
    fecha_nac = db.Column(db.Date)
    creado_en = db.Column(db.DateTime, default=datetime.utcnow)
    activo = db.Column(db.Boolean, default=True)
    citas = db.relationship('Cita', backref='paciente', lazy=True)

class Especialidad(db.Model):
    __tablename__ = 'especialidades'
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(100), nullable=False)
    descripcion = db.Column(db.Text)
    icono = db.Column(db.String(50))
    activa = db.Column(db.Boolean, default=True)
    medicos = db.relationship('Medico', backref='especialidad', lazy=True)


class Medico(db.Model):
    __tablename__ = 'medicos'

    id = db.Column(db.Integer, primary_key=True)

    usuario_id = db.Column(
        db.Integer,
        db.ForeignKey('usuarios.id'),
        unique=True,
        nullable=False
    )

    nombre = db.Column(db.String(150), nullable=False)
    apellido = db.Column(db.String(150), nullable=False)
    cmp = db.Column(db.String(20), unique=True)

    especialidad_id = db.Column(
        db.Integer,
        db.ForeignKey('especialidades.id')
    )

    disponible = db.Column(db.Boolean, default=True)

    citas = db.relationship(
        'Cita',
        backref='medico',
        lazy=True
    )

    horarios = db.relationship(
        'HorarioMedico',
        backref='medico',
        lazy=True
    )

class HorarioMedico(db.Model):
    __tablename__ = 'horarios_medicos'

    id = db.Column(db.Integer, primary_key=True)

    medico_id = db.Column(
        db.Integer,
        db.ForeignKey('medicos.id'),
        nullable=False
    )

    dia_semana = db.Column(
        db.Integer,
        nullable=False
    )

    hora_inicio = db.Column(
        db.Time,
        nullable=False
    )

    hora_fin = db.Column(
        db.Time,
        nullable=False
    )

class Cita(db.Model):
    __tablename__ = 'citas'
    id = db.Column(db.Integer, primary_key=True)
    paciente_id = db.Column(db.Integer, db.ForeignKey('usuarios.id'), nullable=False)
    medico_id = db.Column(db.Integer, db.ForeignKey('medicos.id'), nullable=False)
    especialidad_id = db.Column(db.Integer, db.ForeignKey('especialidades.id'), nullable=False)
    fecha = db.Column(db.Date, nullable=False)
    hora = db.Column(db.String(10), nullable=False)
    motivo = db.Column(db.Text)
    estado = db.Column(db.String(20), default='pendiente')
    tipo = db.Column(db.String(20), default='manual')
    creado_en = db.Column(db.DateTime, default=datetime.utcnow)
    especialidad = db.relationship('Especialidad', backref='citas')

@login_manager.user_loader
def load_user(user_id):
    return Usuario.query.get(int(user_id))

def horario_ocupado(medico_id, fecha, hora, excluir_cita_id=None):
    q = Cita.query.filter_by(medico_id=medico_id, fecha=fecha, hora=hora).filter(Cita.estado != 'cancelada')
    if excluir_cita_id:
        q = q.filter(Cita.id != excluir_cita_id)
    return q.first() is not None

def obtener_datos_reportes():
    por_especialidad = db.session.query(
        Especialidad.nombre, func.count(Cita.id)
    ).join(Cita, Cita.especialidad_id == Especialidad.id).group_by(Especialidad.id).all()

    por_estado = db.session.query(
        Cita.estado, func.count(Cita.id)
    ).group_by(Cita.estado).all()

    hace_6_meses = datetime.utcnow() - timedelta(days=180)
    citas_mes = db.session.query(Cita).filter(Cita.creado_en >= hace_6_meses).all()
    meses_dict = {}
    for c in citas_mes:
        clave = c.creado_en.strftime('%Y-%m')
        meses_dict[clave] = meses_dict.get(clave, 0) + 1
    por_mes = sorted([{'label': k, 'value': v} for k, v in meses_dict.items()])

    top_medicos_raw = db.session.query(
        Medico.nombre, Medico.apellido, func.count(Cita.id).label('total')
    ).join(Cita, Cita.medico_id == Medico.id).group_by(Medico.id).order_by(desc('total')).limit(5).all()

    return {
        'por_especialidad': [{'label': n, 'value': c} for n, c in por_especialidad],
        'por_estado': [{'label': e, 'value': c} for e, c in por_estado],
        'por_mes': por_mes,
        'top_medicos': [{'nombre': f'Dr. {n} {a}', 'citas': t} for n, a, t in top_medicos_raw]
    }
def inicializar_db():
    db.create_all()

    # Migración automática para la columna 'activo'
    try:
        import sqlalchemy as sa
        inspector = sa.inspect(db.engine)
        columns = [c['name'] for c in inspector.get_columns('usuarios')]
        if 'activo' not in columns:
            db.session.execute(sa.text("ALTER TABLE usuarios ADD COLUMN activo TINYINT(1) DEFAULT 1;"))
            db.session.execute(sa.text("UPDATE usuarios SET activo = 1 WHERE activo IS NULL;"))
            db.session.commit()
            print("✅ Columna 'activo' agregada correctamente a la tabla 'usuarios'")
    except Exception as e:
        print("⚠️ Advertencia al aplicar migración 'activo':", e)

    if not Especialidad.query.first():

        esp_list = [
            ('Cardiologia', 'Diagnostico y tratamiento de enfermedades del corazon.'),
            ('Pediatria', 'Atencion medica integral para ninos y adolescentes.'),
            ('Ginecologia', 'Salud reproductiva y atencion a la mujer.'),
            ('Traumatologia', 'Tratamiento de lesiones oseas y musculares.'),
            ('Medicina General', 'Atencion primaria y consulta general.'),
            ('Oftalmologia', 'Diagnostico y tratamiento de enfermedades oculares.'),
            ('Neurologia', 'Enfermedades del sistema nervioso.'),
            ('Dermatologia', 'Diagnostico y tratamiento de enfermedades de la piel.')
        ]

        for n, d in esp_list:
            db.session.add(
                Especialidad(
                    nombre=n,
                    descripcion=d
                )
            )

        db.session.commit()

        medicos_data = [
            ('Carlos', 'Mendoza Ruiz', 'CM-12345', 1),
            ('Ana', 'Garcia Lopez', 'CM-23456', 2),
            ('Rosa', 'Torres Vega', 'CM-34567', 3),
            ('Luis', 'Ramos Flores', 'CM-45678', 4),
            ('Jorge', 'Castillo Pena', 'CM-56789', 5),
        ]

        for n, a, c, e in medicos_data:
            db.session.add(
                Medico(
                    nombre=n,
                    apellido=a,
                    cmp=c,
                    especialidad_id=e
                )
            )

        db.session.add(
            Usuario(
                nombre='Admin',
                apellido='Hospital',
                dni='00000000',
                email='admin@lanfranco.pe',
                password=generate_password_hash('admin123'),
                rol='admin'
            )
        )

        db.session.commit()

        print("DB inicializada OK")

@app.route('/')
def index():
    return render_template('index.html', especialidades=Especialidad.query.filter_by(activa=True).all())

@app.route('/registro', methods=['GET','POST'])
def registro():
    if current_user.is_authenticated: return redirect(url_for('solicitar_cita'))
    if request.method == 'POST':
        nombre=request.form.get('nombre','').strip()
        apellido=request.form.get('apellido','').strip()
        dni=request.form.get('dni','').strip()
        email=request.form.get('email','').strip()
        telefono=request.form.get('telefono','').strip()
        password=request.form.get('password','')
        confirm=request.form.get('confirm_password','')
        fecha_nac_str=request.form.get('fecha_nac','')
        if not all([nombre,apellido,dni,email,password]):
            flash('Todos los campos obligatorios deben completarse.','error'); return render_template('registro.html')
        if password!=confirm:
            flash('Las contrasenas no coinciden.','error'); return render_template('registro.html')
        if len(dni)!=8 or not dni.isdigit():
            flash('El DNI debe tener 8 digitos numericos.','error'); return render_template('registro.html')
        if Usuario.query.filter_by(email=email).first():
            flash('El correo electronico ya esta registrado.','error'); return render_template('registro.html')
        if Usuario.query.filter_by(dni=dni).first():
            flash('El DNI ya esta registrado.','error'); return render_template('registro.html')
        fecha_nac=None
        if fecha_nac_str:
            try: fecha_nac=datetime.strptime(fecha_nac_str,'%Y-%m-%d').date()
            except: pass
        nuevo_usuario = Usuario(nombre=nombre,apellido=apellido,dni=dni,email=email,
            telefono=telefono,password=generate_password_hash(password),fecha_nac=fecha_nac)
        db.session.add(nuevo_usuario)
        db.session.commit()
        
        # Auto-login inmediato para usuarios recién registrados
        login_user(nuevo_usuario)
        
        flash('Registro exitoso! Iniciando sesión automáticamente...','success')
        return redirect(url_for('solicitar_cita'))
    return render_template('registro.html')

@app.route('/login', methods=['GET','POST'])
def login():

    if current_user.is_authenticated:
        return redirect(url_for('solicitar_cita'))

    if request.method == 'POST':

        email = request.form.get('email','').strip()
        password = request.form.get('password','')

        u = Usuario.query.filter_by(email=email).first()

        if u and check_password_hash(u.password,password):
            if not u.activo:
                flash('Cuenta desactivada. Contacta al administrador.', 'error')
                return render_template('login.html')

            login_user(u)

            flash(f'Bienvenido, {u.nombre}!', 'success')

            if u.rol == 'admin':
                return redirect(url_for('admin'))

            elif u.rol == 'doctor':
                return redirect(url_for('doctor_dashboard'))

            else:
                return redirect(url_for('solicitar_cita'))

        flash('Correo o contrasena incorrectos.', 'error')

    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user(); flash('Sesion cerrada correctamente.','info'); return redirect(url_for('index'))

@app.route('/dashboard')
@login_required
def dashboard():
    citas=Cita.query.filter_by(paciente_id=current_user.id).order_by(Cita.fecha.desc()).all()
    return render_template('dashboard.html',citas=citas)

@app.route('/solicitar-cita', methods=['GET','POST'])
def solicitar_cita():
    # Si está autenticado, redirigir admin/doctor a sus paneles
    if current_user.is_authenticated:
        if current_user.rol == 'admin':
            flash('Los administradores no solicitan citas por esta vía.', 'info')
            return redirect(url_for('admin'))
        if current_user.rol == 'doctor':
            flash('Los médicos no solicitan citas por esta vía.', 'info')
            return redirect(url_for('doctor_dashboard'))

    especialidades = Especialidad.query.filter_by(activa=True).all()
    today_str = date.today().isoformat()
    autenticado = current_user.is_authenticated
    nombre_usuario = current_user.nombre if autenticado else ''

    if request.method == 'POST':
        # El POST siempre requiere autenticación como paciente
        if not autenticado:
            flash('Debes iniciar sesión para solicitar una cita.', 'error')
            return redirect(url_for('login'))
        esp_id=request.form.get('especialidad_id')
        med_id=request.form.get('medico_id')
        fecha_s=request.form.get('fecha')
        hora=request.form.get('hora')
        motivo=request.form.get('motivo','')
        tipo=request.form.get('tipo','manual')
        ctx = dict(especialidades=especialidades, today=today_str,
                   autenticado=autenticado, nombre_usuario=nombre_usuario)
        if not all([esp_id,med_id,fecha_s,hora]):
            flash('Completa todos los campos requeridos.','error')
            return render_template('solicitar_cita.html', **ctx)
        try: fecha=datetime.strptime(fecha_s,'%Y-%m-%d').date()
        except:
            flash('Fecha invalida.','error')
            return render_template('solicitar_cita.html', **ctx)
        if fecha<date.today():
            flash('La fecha no puede ser en el pasado.','error')
            return render_template('solicitar_cita.html', **ctx)
        if horario_ocupado(int(med_id), fecha, hora):
            flash('Ese horario ya esta ocupado.','error')
            return render_template('solicitar_cita.html', **ctx)
        db.session.add(Cita(paciente_id=current_user.id,medico_id=int(med_id),
            especialidad_id=int(esp_id),fecha=fecha,hora=hora,motivo=motivo,tipo=tipo))
        db.session.commit()
        cita_nueva = Cita.query.filter_by(
            paciente_id=current_user.id, medico_id=int(med_id),
            fecha=fecha, hora=hora
        ).order_by(Cita.id.desc()).first()
        if cita_nueva and enviar_correo_confirmacion(app, mail, cita_nueva):
            flash('Cita registrada exitosamente! Se envió correo de confirmación.', 'success')
        else:
            flash('Cita registrada exitosamente! No se pudo enviar el correo de confirmación.', 'success')
        return redirect(url_for('dashboard'))

    return render_template('solicitar_cita.html',
        especialidades=especialidades, today=today_str,
        autenticado=autenticado, nombre_usuario=nombre_usuario)


@app.route('/cancelar-cita/<int:cita_id>', methods=['POST'])
@login_required
def cancelar_cita(cita_id):
    c=Cita.query.get_or_404(cita_id)
    if c.paciente_id!=current_user.id and current_user.rol!='admin':
        flash('No tienes permiso.','error'); return redirect(url_for('dashboard'))
    c.estado='cancelada'; db.session.commit()
    flash('Cita cancelada.','info')
    if current_user.rol == 'admin':
        return redirect(url_for('admin'))
    return redirect(url_for('dashboard'))

@app.route('/reprogramar-cita/<int:cita_id>', methods=['POST'])
@login_required
def reprogramar_cita(cita_id):
    c = Cita.query.get_or_404(cita_id)
    if c.paciente_id != current_user.id:
        flash('No tienes permiso.', 'error')
        return redirect(url_for('dashboard'))
    if c.estado not in ['pendiente', 'confirmada', 'reprogramada']:
        flash('No se puede reprogramar esta cita.', 'error')
        return redirect(url_for('dashboard'))
    
    fecha_s = request.form.get('fecha')
    hora = request.form.get('hora')
    
    if not fecha_s or not hora:
        flash('Fecha y hora son requeridas.', 'error')
        return redirect(url_for('dashboard'))
        
    try:
        fecha = datetime.strptime(fecha_s, '%Y-%m-%d').date()
    except:
        flash('Fecha inválida.', 'error')
        return redirect(url_for('dashboard'))
        
    if fecha < date.today():
        flash('La fecha no puede ser en el pasado.', 'error')
        return redirect(url_for('dashboard'))
        
    if horario_ocupado(c.medico_id, fecha, hora, excluir_cita_id=c.id):
        flash('Ese horario ya está ocupado.', 'error')
        return redirect(url_for('dashboard'))
        
    c.fecha = fecha
    c.hora = hora
    c.estado = 'reprogramada'
    db.session.commit()
    
    asunto = f'Reprogramación de cita — {c.especialidad.nombre}'
    if enviar_correo_confirmacion(app, mail, c, asunto=asunto):
        flash('Cita reprogramada exitosamente! Se envió correo de confirmación.', 'success')
    else:
        flash('Cita reprogramada exitosamente! No se pudo enviar el correo de confirmación.', 'success')
        
    return redirect(url_for('dashboard'))

@app.route('/api/clasificar-sintomas', methods=['POST'])
def api_clasificar_sintomas():
    data = request.get_json() or {}
    texto = data.get('texto', '').strip()
    if not texto:
        return jsonify({'error': 'Texto de sintomas vacio'}), 400
    res = clasificar_sintomas(texto)
    return jsonify(res)

@app.route('/api/medicos/<int:esp_id>')
def api_medicos(esp_id):
    m=Medico.query.filter_by(especialidad_id=esp_id,disponible=True).all()
    return jsonify([{'id':x.id,'nombre':f'Dr. {x.nombre} {x.apellido}'} for x in m])

@app.route('/api/horarios/<int:medico_id>/<fecha>')
def api_horarios(medico_id,fecha):
    todos=['08:00','08:30','09:00','09:30','10:00','10:30','11:00','11:30','14:00','14:30','15:00','15:30','16:00','16:30']
    try: fo=datetime.strptime(fecha,'%Y-%m-%d').date()
    except: return jsonify([])
    ocup=[c.hora for c in Cita.query.filter_by(medico_id=medico_id,fecha=fo).filter(Cita.estado!='cancelada').all()]
    return jsonify([h for h in todos if h not in ocup])

@app.route('/admin')
@login_required
def admin():
    if current_user.rol != 'admin':
        flash('Acceso restringido.', 'error')
        return redirect(url_for('dashboard'))
    return render_template(
    'admin.html',
    total_usuarios=Usuario.query.count(),
    total_citas=Cita.query.count(),
    citas_hoy=Cita.query.filter_by(fecha=date.today()).count(),
    citas_pendientes=Cita.query.filter_by(estado='pendiente').count(),
    citas_recientes=Cita.query.order_by(Cita.creado_en.desc()).limit(15).all(),
    pacientes=Usuario.query.filter_by(rol='paciente').all(),
    medicos=Medico.query.all(),
    especialidades=Especialidad.query.all(),
    horarios=HorarioMedico.query.all(),
    datos_reportes=obtener_datos_reportes()
)    
@app.route('/admin/confirmar/<int:cita_id>', methods=['POST'])
@login_required
def confirmar_cita(cita_id):
    if current_user.rol!='admin': return jsonify({'error':'No autorizado'}),403
    c=Cita.query.get_or_404(cita_id); c.estado='confirmada'; db.session.commit()
    enviar_correo_confirmacion(app, mail, c, asunto='Su cita ha sido confirmada')
    return jsonify({'ok':True})
@app.route('/admin/toggle-medico/<int:medico_id>', methods=['POST'])
@login_required
def toggle_medico(medico_id):
    if current_user.rol != 'admin':
        return jsonify({'error': 'No autorizado'}), 403
    m = Medico.query.get_or_404(medico_id)
    m.disponible = not m.disponible
    db.session.commit()
    return jsonify({'ok': True, 'disponible': m.disponible})

@app.route('/admin/cita/editar/<int:cita_id>', methods=['POST'])
@login_required
def editar_cita_admin(cita_id):
    if current_user.rol != 'admin':
        flash('Acceso restringido.', 'error')
        return redirect(url_for('dashboard'))
    
    c = Cita.query.get_or_404(cita_id)
    medico_id = request.form.get('medico_id')
    fecha_str = request.form.get('fecha')
    hora = request.form.get('hora')
    
    if not all([medico_id, fecha_str, hora]):
        flash('Todos los campos son requeridos.', 'error')
        return redirect(url_for('admin'))
        
    try:
        fecha = datetime.strptime(fecha_str, '%Y-%m-%d').date()
    except:
        flash('Fecha inválida.', 'error')
        return redirect(url_for('admin'))
        
    if fecha < date.today():
        flash('La fecha no puede ser en el pasado.', 'error')
        return redirect(url_for('admin'))
        
    if horario_ocupado(int(medico_id), fecha, hora, excluir_cita_id=c.id):
        flash('Ese horario ya está ocupado.', 'error')
        return redirect(url_for('admin'))
        
    c.medico_id = int(medico_id)
    c.fecha = fecha
    c.hora = hora
    db.session.commit()
    
    flash('Cita modificada correctamente.', 'success')
    return redirect(url_for('admin'))

@app.route('/admin/usuario/editar/<int:user_id>', methods=['POST'])
@login_required
def editar_usuario(user_id):
    if current_user.rol != 'admin':
        flash('Acceso restringido.', 'error')
        return redirect(url_for('dashboard'))
    
    u = Usuario.query.get_or_404(user_id)
    nombre = request.form.get('nombre', '').strip()
    apellido = request.form.get('apellido', '').strip()
    email = request.form.get('email', '').strip()
    telefono = request.form.get('telefono', '').strip()
    
    if not all([nombre, apellido, email]):
        flash('Nombre, apellido y correo son obligatorios.', 'error')
        return redirect(url_for('admin'))
        
    # Validar que el email no esté tomado por otro usuario
    existing = Usuario.query.filter(Usuario.email == email, Usuario.id != user_id).first()
    if existing:
        flash('El correo electrónico ya está registrado por otro usuario.', 'error')
        return redirect(url_for('admin'))
        
    u.nombre = nombre
    u.apellido = apellido
    u.email = email
    u.telefono = telefono
    db.session.commit()
    
    if u.rol == 'doctor':
        m = Medico.query.filter_by(usuario_id=u.id).first()
        if m:
            m.nombre = nombre
            m.apellido = apellido
            db.session.commit()
            
    flash('Usuario actualizado correctamente.', 'success')
    return redirect(url_for('admin'))

@app.route('/admin/usuario/toggle/<int:user_id>', methods=['POST'])
@login_required
def toggle_usuario(user_id):
    if current_user.rol != 'admin':
        return jsonify({'error': 'No autorizado'}), 403
    
    u = Usuario.query.get_or_404(user_id)
    if u.id == current_user.id:
        return jsonify({'error': 'No puedes desactivar tu propia cuenta'}), 400
        
    u.activo = not u.activo
    db.session.commit()
    return jsonify({'ok': True, 'activo': u.activo})
    
@app.route('/admin/doctor/crear', methods=['POST'])
@login_required
def crear_doctor():

    if current_user.rol != 'admin':
        return redirect(url_for('dashboard'))

    if Usuario.query.filter_by(dni=request.form['dni']).first():
        flash('El DNI ya está registrado.', 'error')
        return redirect(url_for('admin'))

    if Usuario.query.filter_by(email=request.form['email']).first():
        flash('El correo ya está registrado.', 'error')
        return redirect(url_for('admin'))

    usuario = Usuario(
        nombre=request.form['nombre'],
        apellido=request.form['apellido'],
        dni=request.form['dni'],
        email=request.form['email'],
        telefono=request.form['telefono'],
        password=generate_password_hash(request.form['password']),
        rol='doctor'
    )

    db.session.add(usuario)
    db.session.commit()

    medico = Medico(
        usuario_id=usuario.id,
        nombre=request.form['nombre'],
        apellido=request.form['apellido'],
        cmp=request.form['cmp'],
        especialidad_id=request.form['especialidad_id']
    )

    db.session.add(medico)
    db.session.commit()

    flash('Doctor registrado correctamente', 'success')

    return redirect(url_for('admin'))

@app.route('/doctor')
@login_required
def doctor_dashboard():

    if current_user.rol != 'doctor':
        return redirect(url_for('dashboard'))

    medico = Medico.query.filter_by(
        usuario_id=current_user.id
    ).first()

    if not medico:
        flash(
            'No se encontró el médico asociado a esta cuenta.',
            'error'
        )
        return redirect(url_for('logout'))

    citas = Cita.query.filter_by(
        medico_id=medico.id
    ).order_by(
        Cita.fecha.asc(),
        Cita.hora.asc()
    ).all()

    for cita in citas:
        cita.dia_semana = cita.fecha.weekday()

    return render_template(
        'doctor.html',
        medico=medico,
        citas=citas
    )

@app.route('/admin/especialidad/crear', methods=['POST'])
@login_required
def crear_especialidad():

    if current_user.rol != 'admin':
        return redirect(url_for('dashboard'))

    nombre = request.form['nombre'].strip()
    descripcion = request.form.get('descripcion', '').strip()

    if Especialidad.query.filter_by(nombre=nombre).first():
        flash('La especialidad ya existe.', 'error')
        return redirect(url_for('admin'))

    db.session.add(
        Especialidad(
            nombre=nombre,
            descripcion=descripcion,
            activa=True
        )
    )

    db.session.commit()

    flash('Especialidad registrada correctamente.', 'success')

    return redirect(url_for('admin'))
@app.route('/admin/especialidad/editar/<int:id>', methods=['POST'])
@login_required
def editar_especialidad(id):

    if current_user.rol != 'admin':
        return redirect(url_for('dashboard'))

    esp = Especialidad.query.get_or_404(id)

    esp.nombre = request.form['nombre'].strip()
    esp.descripcion = request.form.get('descripcion', '').strip()

    db.session.commit()

    flash('Especialidad actualizada correctamente.', 'success')

    return redirect(url_for('admin'))

@app.route('/admin/toggle-especialidad/<int:id>', methods=['POST'])
@login_required
def toggle_especialidad(id):

    if current_user.rol != 'admin':
        return jsonify({'error':'No autorizado'}), 403

    esp = Especialidad.query.get_or_404(id)

    esp.activa = not esp.activa

    db.session.commit()

    return jsonify({
        'ok': True,
        'activa': esp.activa
    })

@app.route('/admin/horario/crear', methods=['POST'])
@login_required
def crear_horario():

    if current_user.rol != 'admin':
        return redirect(url_for('dashboard'))

    horario = HorarioMedico(
        medico_id=int(request.form['medico_id']),
        dia_semana=int(request.form['dia_semana']),
        hora_inicio=datetime.strptime(
            request.form['hora_inicio'],
            '%H:%M'
        ).time(),
        hora_fin=datetime.strptime(
            request.form['hora_fin'],
            '%H:%M'
        ).time()
    )

    db.session.add(horario)
    db.session.commit()

    flash(
        'Horario registrado correctamente.',
        'success'
    )

    return redirect(url_for('admin'))

@app.route(
    '/admin/horario/eliminar/<int:id>',
    methods=['POST']
)
@login_required
def eliminar_horario(id):

    if current_user.rol != 'admin':
        return redirect(url_for('dashboard'))

    horario = HorarioMedico.query.get_or_404(id)

    db.session.delete(horario)

    db.session.commit()

    flash(
        'Horario eliminado correctamente',
        'success'
    )

    return redirect(url_for('admin'))
@app.route(
    '/admin/horario/editar/<int:id>',
    methods=['POST']
)
@login_required
def editar_horario(id):

    if current_user.rol != 'admin':
        return redirect(url_for('dashboard'))

    horario = HorarioMedico.query.get_or_404(id)

    horario.dia_semana = int(
        request.form['dia_semana']
    )

    horario.hora_inicio = datetime.strptime(
        request.form['hora_inicio'],
        '%H:%M'
    ).time()

    horario.hora_fin = datetime.strptime(
        request.form['hora_fin'],
        '%H:%M'
    ).time()

    db.session.commit()

    flash(
        'Horario actualizado correctamente',
        'success'
    )

    return redirect(url_for('admin'))

@app.route('/admin/reportes/exportar-pdf')
@login_required
def exportar_reporte_pdf():
    if current_user.rol != 'admin':
        flash('Acceso restringido.', 'error')
        return redirect(url_for('dashboard'))

    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    datos = obtener_datos_reportes()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph('Reporte de Citas — Hospital Carlos Lanfranco La Hoz', styles['Title']),
        Spacer(1, 16)
    ]

    def tabla_titulo(texto):
        elements.append(Paragraph(texto, styles['Heading2']))
        elements.append(Spacer(1, 8))

    def crear_tabla(headers, rows):
        data = [headers] + rows
        t = Table(data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a4b8c')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f4f7fb')])
        ]))
        elements.append(t)
        elements.append(Spacer(1, 20))

    tabla_titulo('Citas por Especialidad')
    crear_tabla(['Especialidad', 'Cantidad'], [[d['label'], str(d['value'])] for d in datos['por_especialidad']])

    tabla_titulo('Citas por Estado')
    crear_tabla(['Estado', 'Cantidad'], [[d['label'], str(d['value'])] for d in datos['por_estado']])

    tabla_titulo('Top 5 Médicos')
    crear_tabla(['Médico', 'Citas'], [[d['nombre'], str(d['citas'])] for d in datos['top_medicos']])

    doc.build(elements)
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True,
                     download_name='reporte_citas.pdf')

@app.route('/admin/reportes/exportar-excel')
@login_required
def exportar_reporte_excel():
    if current_user.rol != 'admin':
        flash('Acceso restringido.', 'error')
        return redirect(url_for('dashboard'))

    from openpyxl import Workbook

    datos = obtener_datos_reportes()
    wb = Workbook()

    ws1 = wb.active
    ws1.title = 'Especialidad'
    ws1.append(['Especialidad', 'Cantidad'])
    for d in datos['por_especialidad']:
        ws1.append([d['label'], d['value']])

    ws2 = wb.create_sheet('Estado')
    ws2.append(['Estado', 'Cantidad'])
    for d in datos['por_estado']:
        ws2.append([d['label'], d['value']])

    ws3 = wb.create_sheet('PorMes')
    ws3.append(['Mes', 'Cantidad'])
    for d in datos['por_mes']:
        ws3.append([d['label'], d['value']])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(buffer,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='reporte_citas.xlsx')

if __name__=='__main__':
    with app.app_context():
        inicializar_db()
    app.run(debug=True,port=5000)


